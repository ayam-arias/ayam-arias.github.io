"""Agrega los certificados nuevos a los documentos madre de 02_INFORMACION PERSONAL:
el CV base (Word), el Currículum Maestro de Capacitaciones (Word) y su Excel.

Un PDF se agrega una sola vez: su hash queda en scripts/documentos_vistos.txt.
La primera ejecución solo registra los PDF actuales (los documentos ya están completos).
"""
import copy, datetime, os, re, shutil, unicodedata
from pathlib import Path

DOCS = Path(os.environ.get("INFO_PERSONAL", Path.home() / "OneDrive/Documentos/02_INFORMACION PERSONAL"))
CV_BASE = DOCS / "CV_MAESTRO_Ian_Arias_Carrasco_2026.docx"
COMPILADO = DOCS / "Curriculum_Maestro_Capacitaciones_Ian_Arias.docx"
EXCEL = DOCS / "Curriculum_Maestro_Capacitaciones_Ian_Arias.xlsx"
VISTOS = Path(__file__).resolve().parent / "documentos_vistos.txt"
RAYA = "—"
MESES = "enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre".split()
# Encabezados del CV base que no se llaman igual que la carpeta.
ALIAS_CV = {"Ciencia de Datos e IA": "Ciencia de Datos e Inteligencia Artificial"}


def norm(s):
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower().strip()


def abierto(p):
    """Word/Excel crean '~$' + nombre mientras el archivo está abierto."""
    return (p.parent / ("~$" + p.name[2:])).exists() or (p.parent / ("~$" + p.name)).exists()


def poner_texto(parrafo, *textos):
    """Reemplaza el texto conservando el formato de las primeras runs."""
    runs = parrafo.runs
    for i, r in enumerate(runs):
        r.text = textos[i] if i < len(textos) else ""
    if len(textos) > len(runs):
        runs[-1].text += "".join(textos[len(runs):])


def sumar_en_texto(doc_parrafos, patron, n):
    """Suma n al número que precede a `patron` (ej. '175 capacitaciones certificadas')."""
    rx = re.compile(r"(\d+)(\s+" + patron + ")")
    for p in doc_parrafos:
        for r in p.runs:
            if rx.search(r.text):
                r.text = rx.sub(lambda m: f"{int(m.group(1)) + n}{m.group(2)}", r.text)


def fila_tabla(tabla, valores):
    tr = copy.deepcopy(tabla.rows[-1]._tr)
    tabla._tbl.append(tr)
    for celda, v in zip(tabla.rows[-1].cells, valores):
        poner_texto(celda.paragraphs[0], str(v))


def actualizar_compilado(items, hoy):
    import docx
    d = docx.Document(COMPILADO)
    tablas, faltan, agregados = {}, [], 0
    # Cada sección es un Heading 1 "NN. Categoría" seguido de su tabla.
    body = list(d.element.body.iterchildren())
    for i, el in enumerate(body):
        if el.tag.endswith("}p") and "Heading1" in el.xml:
            m = re.match(r"(\d\d)\.", docx.text.paragraph.Paragraph(el, d._body).text)
            nxt = next((b for b in body[i + 1:] if b.tag.endswith("}tbl")), None)
            if m and nxt is not None:
                tablas[m.group(1)] = docx.table.Table(nxt, d._body)
    resumen = d.tables[0]
    for it in items:
        for carpeta in it["carpetas"]:
            num = carpeta[:2]
            if num not in tablas:
                faltan.append(carpeta)
                continue
            fila_tabla(tablas[num], [it["inst"] or RAYA, it["titulo"], it["horas"] or RAYA, it["fecha"]])
            agregados += 1
            for row in resumen.rows:
                if row.cells[1].text.startswith(num + "."):
                    poner_texto(row.cells[2].paragraphs[0], str(int(row.cells[2].text) + 1))
            # "24 capacitación(es) registrada(s)" entre el título de la sección y su tabla
            el = tablas[num]._tbl.getprevious()
            while el is not None and "Heading1" not in el.xml:
                if "registrada" in "".join(el.itertext()):
                    sumar_en_texto([docx.text.paragraph.Paragraph(el, d._body)], "capacitaci", 1)
                    break
                el = el.getprevious()
    for row in resumen.rows:
        if row.cells[1].text.strip() == "TOTAL":
            poner_texto(row.cells[2].paragraphs[0], str(int(row.cells[2].text) + agregados))
    sumar_en_texto(d.paragraphs, "capacitaciones registradas", agregados)
    for p in d.paragraphs:
        for r in p.runs:
            r.text = re.sub(r"Actualizado en \w+ de \d{4}", f"Actualizado en {MESES[hoy.month - 1]} de {hoy.year}", r.text)
    d.save(COMPILADO)
    return agregados, faltan


def actualizar_cv_base(items):
    import docx
    d = docx.Document(CV_BASE)
    ps = d.paragraphs
    ini = next(i for i, p in enumerate(ps) if p.text.strip() == "CAPACITACIONES")
    encabezados = {norm(p.text): p for p in ps[ini + 1:] if p.text.strip() and p.style.name != "List Paragraph"}
    faltan, agregados = [], 0
    for it in items:
        linea = (it["titulo"] + (f" ({it['horas']} h)" if it["horas"] else ""), (f", {it['inst']}" if it["inst"] else "") + f" ({it['fecha'][-4:]}).")
        destinos = [ALIAS_CV.get(a, a) for a in it["areas"] if a != "Títulos y Diplomados"]
        if it["tipo"] == "diplomado":
            destinos.insert(0, "Diplomados")
        for a in destinos:
            h = encabezados.get(norm(a))
            if h is None:
                faltan.append(a)
                continue
            modelo = h._p.getnext()  # primera viñeta de la sección: se copia su formato
            nuevo = copy.deepcopy(modelo)
            h._p.addnext(nuevo)
            poner_texto(docx.text.paragraph.Paragraph(nuevo, h._parent), *linea)
            agregados += 1
    sumar_en_texto(d.paragraphs, "capacitaciones certificadas", len(items))
    d.save(CV_BASE)
    return agregados, faltan


def actualizar_excel(items):
    import docx, openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    if EXCEL.exists():
        wb = openpyxl.load_workbook(EXCEL)
        ws = wb.active
    else:  # se crea una vez a partir de las tablas del compilado
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Capacitaciones"
        ws.append(["Categoría", "Institución", "Curso / Capacitación", "Duración (h)", "Fecha", "Archivo PDF"])
        d = docx.Document(COMPILADO)
        cat = None
        for el in d.element.body.iterchildren():
            if el.tag.endswith("}p") and "Heading1" in el.xml:
                cat = docx.text.paragraph.Paragraph(el, d._body).text.strip()
            elif el.tag.endswith("}tbl") and cat and re.match(r"\d\d\.", cat):
                for row in docx.table.Table(el, d._body).rows[1:]:
                    inst, curso, dur, fecha = (c.text.strip() for c in row.cells[:4])
                    ws.append([cat, inst, curso, int(dur) if dur.isdigit() else None, None if fecha == RAYA else fecha, None])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1B2A4A")
        for col, w in zip("ABCDEF", (44, 40, 80, 13, 22, 60)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
    cats = {str(r[0].value)[:2]: r[0].value for r in ws.iter_rows(min_row=2) if r[0].value}
    for it in items:
        for carpeta in it["carpetas"]:
            ws.append([cats.get(carpeta[:2], re.sub(r"^(\d\d) ", r"\1. ", carpeta)), it["inst"] or None, it["titulo"], it["horas"], it["fecha"], it["ruta"]])
    ref = f"A1:F{ws.max_row}"
    if "Capacitaciones" in ws.tables:
        ws.tables["Capacitaciones"].ref = ref
    else:
        t = Table(displayName="Capacitaciones", ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    wb.save(EXCEL)
    return ws.max_row - 1


def actualizar(fuentes, describir, dry=False):
    """fuentes: {hash: [rutas]}; describir(rutas) -> dict del certificado."""
    if not VISTOS.exists():
        if not dry:
            VISTOS.write_text("\n".join(sorted(fuentes)) + "\n", encoding="utf-8")
        print("Documentos: primera ejecución, se registraron los PDF actuales como ya incluidos.")
        if not dry and not EXCEL.exists():
            print(f"Documentos: Excel creado con {actualizar_excel([])} filas.")
        return
    vistos = set(VISTOS.read_text(encoding="utf-8").split())
    nuevos = [h for h in fuentes if h not in vistos]
    if not nuevos:
        print("Documentos: sin cambios.")
        return
    items = [describir(fuentes[h]) for h in nuevos]
    for it in items:
        print(f"Documentos + {it['titulo']} | {it['fecha']}")
    if dry:
        return
    ocupados = [p.name for p in (CV_BASE, COMPILADO, EXCEL) if abierto(p)]
    if ocupados:  # no se marcan como vistos: se reintenta la próxima vez
        print(f"Documentos: se omiten porque están abiertos ({', '.join(ocupados)}). Ciérralos y vuelve a ejecutar.")
        return
    hoy = datetime.date.today()
    respaldo = DOCS / "_respaldos"
    respaldo.mkdir(exist_ok=True)
    for p in (CV_BASE, COMPILADO, EXCEL):
        if p.exists():
            shutil.copy2(p, respaldo / f"{p.stem}_{hoy:%Y-%m-%d}{p.suffix}")
    filas = actualizar_excel(items)  # primero: si no existe, se crea desde el compilado sin modificar
    n1, f1 = actualizar_compilado(items, hoy)
    n2, f2 = actualizar_cv_base(items)
    with VISTOS.open("a", encoding="utf-8") as f:
        f.write("\n".join(nuevos) + "\n")
    print(f"Documentos: {len(items)} certificado(s) agregados (compilado {n1} filas, CV base {n2} viñetas, Excel {filas} filas).")
    for falta in sorted(set(f1 + f2)):
        print(f"  Aviso: no existe la sección «{falta}» en algún documento; agrégala a mano.")
